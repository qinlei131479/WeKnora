import asyncio
import csv
import io
import json
import os
import uuid
from pathlib import Path
from collections.abc import Iterable
from typing import Any
from urllib.parse import quote

from fastapi import HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from yuxi.agents.backends.sandbox import (
    ensure_thread_dirs,
    sandbox_uploads_dir,
)
from yuxi.config import config as app_config
from yuxi.repositories.agent_repository import AgentRepository
from yuxi.repositories.conversation_repository import INVOCATION_CONVERSATION_SOURCES, ConversationRepository
from yuxi.services.mention_search_service import invalidate_mention_cache
from yuxi.services.ocr_service import parse_document
from yuxi.storage.minio import StorageError, get_minio_client
from yuxi.storage.postgres.models_business import AgentRun, User
from yuxi.utils.datetime_utils import utc_isoformat
from yuxi.utils.logging_config import logger
from yuxi.utils.paths import VIRTUAL_PATH_UPLOADS
from yuxi.utils.upload_utils import read_upload_with_limit

ATTACHMENT_ALLOWED_EXTENSIONS: tuple[str, ...] = ()
MAX_ATTACHMENT_MARKDOWN_CHARS = 32_000  # TODO: 转 MARKDOWN的时候，不应该裁剪
DEFAULT_ATTACHMENT_MAX_SIZE_MB = 10
TMP_ATTACHMENT_PREFIX = "tmp/chat_attachments"
TMP_ATTACHMENT_DOCUMENT_PARSE_EXTENSIONS = (
    ".pdf",
    ".doc",
    ".docx",
    ".ppt",
    ".pptx",
    ".rtf",
    ".html",
    ".htm",
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tiff",
    ".tif",
)
TMP_ATTACHMENT_TEXT_EXTENSIONS = (".txt", ".md", ".markdown")
TMP_ATTACHMENT_TABLE_EXTENSIONS = (".xlsx", ".xls", ".csv")
TMP_ATTACHMENT_PARSE_EXTENSIONS = (
    *TMP_ATTACHMENT_DOCUMENT_PARSE_EXTENSIONS,
    *TMP_ATTACHMENT_TEXT_EXTENSIONS,
    *TMP_ATTACHMENT_TABLE_EXTENSIONS,
)
TMP_ATTACHMENT_IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif")
TMP_ATTACHMENT_OCR_METHODS = ("weknora",)
TMP_ATTACHMENT_PARSE_METHODS = ("weknora",)
TMP_ATTACHMENT_TABLE_MAX_SHEETS = 20
TMP_ATTACHMENT_TABLE_SAMPLE_ROWS = 20
TMP_ATTACHMENT_TABLE_MAX_COLUMNS = 60
TMP_ATTACHMENT_TABLE_MAX_ROWS = 10_000
TMP_ATTACHMENT_TABLE_MAX_CELLS = 120_000
# 工作线程无法被 wait_for 真正取消，超时只是请求侧兜底；真正的边界由行数/单元格预算保证。
TMP_ATTACHMENT_TABLE_SECONDS = 15.0


def _truncate_markdown(markdown: str) -> tuple[str, bool]:
    if len(markdown) <= MAX_ATTACHMENT_MARKDOWN_CHARS:
        return markdown, False

    truncated_content = markdown[: MAX_ATTACHMENT_MARKDOWN_CHARS - 100].rstrip()
    truncated_content = f"{truncated_content}\n\n[内容已截断，超出 {MAX_ATTACHMENT_MARKDOWN_CHARS} 字符限制]"
    return truncated_content, True


def _get_attachment_max_size_bytes() -> int:
    """读取附件上限配置，非法配置回退到 10 MB。"""
    raw_value = os.getenv("AGENT_ATTACHMENT_MAX_SIZE_MB", str(DEFAULT_ATTACHMENT_MAX_SIZE_MB))
    try:
        max_size_mb = int(raw_value)
    except ValueError:
        max_size_mb = 0
    if max_size_mb <= 0:
        logger.warning(
            f"Invalid AGENT_ATTACHMENT_MAX_SIZE_MB={raw_value}, fallback to {DEFAULT_ATTACHMENT_MAX_SIZE_MB} MB"
        )
        max_size_mb = DEFAULT_ATTACHMENT_MAX_SIZE_MB
    return max_size_mb * 1024 * 1024


async def require_user_conversation(conv_repo: ConversationRepository, thread_id: str, uid: str):
    conversation = await conv_repo.get_conversation_by_thread_id(thread_id)
    if not conversation or conversation.uid != str(uid) or conversation.status == "deleted":
        raise HTTPException(status_code=404, detail="对话线程不存在")
    return conversation


def _safe_file_name(file_name: str | None, default: str = "attachment.bin") -> str:
    safe_name = Path(file_name or "").name.replace("/", "_").replace("\\", "_").strip(" .")
    return safe_name or default


def _make_upload_virtual_path(file_name: str) -> str:
    return f"{VIRTUAL_PATH_UPLOADS}/{_safe_file_name(file_name)}"


def _make_attachment_path(file_name: str) -> str:
    """生成附件在沙盒用户目录中的统一路径。"""
    file_name = _safe_file_name(file_name)
    # 提取不带扩展名的部分
    base_name = file_name
    for ext in [".docx", ".txt", ".html", ".htm", ".pdf", ".md"]:
        if file_name.lower().endswith(ext):
            base_name = file_name[: -len(ext)]
            break

    # 替换路径分隔符
    safe_name = base_name.replace("/", "_").replace("\\", "_")
    return f"{safe_name}.md"


def _build_attachment_storage_path(*, uid: str, thread_id: str, file_name: str) -> tuple[str, Path]:
    """返回附件虚拟路径和宿主机落盘路径。"""
    relative_name = _make_attachment_path(file_name)
    virtual_path = f"{VIRTUAL_PATH_UPLOADS}/attachments/{relative_name}"

    host_dir = Path(app_config.save_dir) / "threads" / thread_id / "user-data" / "uploads" / "attachments"
    host_dir.mkdir(parents=True, exist_ok=True)
    host_path = host_dir / relative_name

    return virtual_path, host_path


def _artifact_url(thread_id: str, virtual_path: str) -> str:
    return f"/api/chat/thread/{thread_id}/artifacts/{virtual_path.lstrip('/')}"


def _tmp_attachment_prefix(uid: str, tmp_file_id: str) -> str:
    return f"{TMP_ATTACHMENT_PREFIX}/{uid}/{tmp_file_id}"


def _get_tmp_attachment_bucket() -> str:
    return get_minio_client().KB_BUCKETS["documents"]


def _is_allowed_tmp_attachment_bucket(bucket_name: str, expected_bucket: str) -> bool:
    """Allow the pre-split bucket only for existing attachment references."""
    return bucket_name in {expected_bucket, "knowledgebases"}


def _make_tmp_attachment_object(uid: str, file_name: str) -> tuple[str, str]:
    """生成用户隔离的 tmp 对象路径。"""
    tmp_file_id = uuid.uuid4().hex
    safe_name = _safe_file_name(file_name)
    return tmp_file_id, f"{_tmp_attachment_prefix(uid, tmp_file_id)}/original/{safe_name}"


def _make_tmp_parsed_object(uid: str, tmp_file_id: str, file_name: str) -> str:
    stem = Path(_safe_file_name(file_name)).stem or "attachment"
    return f"{_tmp_attachment_prefix(uid, tmp_file_id)}/parsed/{stem}.md"


def _make_tmp_table_index_object(uid: str, tmp_file_id: str, file_name: str) -> str:
    stem = Path(_safe_file_name(file_name)).stem or "attachment"
    return f"{_tmp_attachment_prefix(uid, tmp_file_id)}/parsed/{stem}.index.json"


def _minio_source(bucket_name: str, object_name: str) -> str:
    return f"minio://{bucket_name}/{quote(object_name, safe='/')}"


def _parse_user_tmp_object(object_name: str, uid: str) -> tuple[str, str, str]:
    if not object_name or "\\" in object_name:
        raise HTTPException(status_code=400, detail="无效的临时附件路径")

    user_prefix = f"{TMP_ATTACHMENT_PREFIX}/{uid}/"
    if not object_name.startswith(user_prefix):
        raise HTTPException(status_code=403, detail="无权访问该临时附件")

    parts = object_name[len(user_prefix) :].split("/")
    if len(parts) != 3 or any(not part or part in {".", ".."} for part in parts):
        raise HTTPException(status_code=400, detail="无效的临时附件路径")

    return parts[0], parts[1], parts[2]


def _require_tmp_object_section(
    object_name: str,
    uid: str,
    section: str,
    tmp_file_id: str | None = None,
) -> tuple[str, str]:
    current_tmp_file_id, current_section, object_file_name = _parse_user_tmp_object(object_name, uid)
    if current_section != section or (tmp_file_id is not None and current_tmp_file_id != tmp_file_id):
        raise HTTPException(status_code=400, detail="无效的临时附件路径")
    if section == "parsed" and Path(object_file_name).suffix.lower() != ".md":
        raise HTTPException(status_code=400, detail="无效的解析附件路径")
    return current_tmp_file_id, object_file_name


def _normalize_parse_method(file_name: str, parse_method: str | None) -> str:
    """按文件类型确定临时附件解析方式。"""

    suffix = Path(file_name).suffix.lower()
    if suffix not in TMP_ATTACHMENT_PARSE_EXTENSIONS:
        raise HTTPException(status_code=400, detail="当前附件类型不支持解析")

    method = parse_method or "weknora"
    allowed_methods = TMP_ATTACHMENT_PARSE_METHODS

    if method not in allowed_methods:
        allowed = ", ".join(allowed_methods)
        raise HTTPException(status_code=400, detail=f"不支持的解析方法: {method}，可选: {allowed}")
    return method


def _clean_cell_value(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", " ").replace("\n", " ").strip()
    return text


def _build_table_markdown(index: dict) -> str:
    sections = [f"# {index['file_name']} 表格索引", "", "该文件是结构化表格；精确行级取数请读取原始文件。", ""]
    for sheet in index.get("sheets", []):
        sections.extend([f"## Sheet：{sheet['name']}", ""])
        sections.append(
            f"表头行：{sheet['header_row']}；有效行：{sheet['first_data_row']}-{sheet['last_data_row']}；"
            f"总行数：{sheet['row_count']}；列数：{sheet['column_count']}"
        )
        sections.append("")
        headers = sheet.get("headers") or []
        if headers:
            sections.append("| " + " | ".join(str(item) for item in headers) + " |")
            sections.append("|" + "---|" * len(headers))
            for row in sheet.get("sample_rows", []):
                cells = row.get("cells") or []
                cells = cells + [""] * (len(headers) - len(cells))
                sections.append("| " + " | ".join(_clean_cell_value(cell) for cell in cells[: len(headers)]) + " |")
        if sheet.get("truncated"):
            sections.append("")
            sections.append("> 索引仅保留采样行；不是完整数据源。")
        sections.append("")
    markdown = "\n".join(sections)
    markdown, truncated = _truncate_markdown(markdown)
    index["preview_truncated"] = truncated
    return markdown


def _bounded_table_rows(rows: Iterable[tuple[str, ...]]) -> list[tuple[str, ...]]:
    """边读边校验预算，超限立即中止，避免把超大工作簿整体读进内存。"""
    collected: list[tuple[str, ...]] = []
    cell_count = 0
    for row in rows:
        collected.append(row)
        cell_count += len(row)
        if len(collected) > TMP_ATTACHMENT_TABLE_MAX_ROWS:
            raise ValueError(f"表格行数超过 {TMP_ATTACHMENT_TABLE_MAX_ROWS} 行限制")
        if cell_count > TMP_ATTACHMENT_TABLE_MAX_CELLS:
            raise ValueError(f"表格单元格数超过 {TMP_ATTACHMENT_TABLE_MAX_CELLS} 限制")
    return collected


def _empty_sheet_index(*, name: str, sheet_index: int) -> dict:
    return {
        "name": name,
        "index": sheet_index,
        "header_row": 0,
        "headers": [],
        "first_data_row": 1,
        "last_data_row": 0,
        "row_count": 0,
        "column_count": 0,
        "sample_rows": [],
        "truncated": False,
    }


def _build_sheet_index(
    *,
    name: str,
    sheet_index: int,
    rows: list[tuple[str, ...]],
    header_policy: str = "auto",
    allow_empty: bool = False,
) -> dict:
    """基于物理行构建索引；过滤空行时保留原始行号。

    header_policy: auto 取首个多列非空行（工作簿常有标题行）；first 固定首行（CSV 无合并标题）。
    """
    physical_rows = [
        (row_number, row)
        for row_number, row in enumerate(rows, start=1)
        if any(row)
    ]
    if not physical_rows:
        if allow_empty:
            return _empty_sheet_index(name=name, sheet_index=sheet_index)
        raise ValueError("表格中没有有效数据行")

    if header_policy == "first":
        header_position = 0
    else:
        header_position = next(
            (
                position
                for position, (_row_number, row) in enumerate(physical_rows)
                if sum(bool(cell) for cell in row) >= 2
            ),
            0,
        )
    column_count = max((len(row) for _row_number, row in physical_rows), default=0)
    header_row_number = physical_rows[header_position][0]
    headers = list(physical_rows[header_position][1][:TMP_ATTACHMENT_TABLE_MAX_COLUMNS])
    data_positions = range(header_position + 1, len(physical_rows))
    sample_rows = [
        {
            "row": physical_rows[position][0],
            "cells": list(physical_rows[position][1][:TMP_ATTACHMENT_TABLE_MAX_COLUMNS]),
        }
        for position in data_positions
    ][:TMP_ATTACHMENT_TABLE_SAMPLE_ROWS]
    next_data_position = header_position + 1
    first_data_row = (
        physical_rows[next_data_position][0]
        if next_data_position < len(physical_rows)
        else header_row_number
    )
    last_data_row = physical_rows[-1][0]
    row_count = max(len(physical_rows) - header_position - 1, 0)
    return {
        "name": name,
        "index": sheet_index,
        "header_row": header_row_number,
        "headers": headers,
        "first_data_row": first_data_row,
        "last_data_row": last_data_row,
        "row_count": row_count,
        "column_count": min(column_count, TMP_ATTACHMENT_TABLE_MAX_COLUMNS),
        "sample_rows": sample_rows,
        "truncated": row_count > TMP_ATTACHMENT_TABLE_SAMPLE_ROWS,
    }


def _extract_xlsx_index(file_name: str, file_content: bytes) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
    sheets: list[dict] = []
    try:
        total_sheets = len(workbook.worksheets)
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            try:
                rows = _bounded_table_rows(
                    tuple(_clean_cell_value(cell) for cell in row)
                    for row in worksheet.iter_rows(values_only=True)
                )
                sheets.append(
                    _build_sheet_index(
                        name=worksheet.title,
                        sheet_index=sheet_index,
                        rows=rows,
                    )
                )
            except ValueError:
                continue
            if len(sheets) >= TMP_ATTACHMENT_TABLE_MAX_SHEETS:
                break
    finally:
        workbook.close()
    if not sheets:
        raise ValueError("表格中没有有效数据行")
    return {
        "file_name": file_name,
        "file_kind": "spreadsheet",
        "sheets": sheets,
        "sheets_truncated": total_sheets > len(sheets),
    }


def _extract_xls_index(file_name: str, file_content: bytes) -> dict:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=file_content)
    sheets: list[dict] = []
    for sheet_index, worksheet in enumerate(workbook.sheets()[:TMP_ATTACHMENT_TABLE_MAX_SHEETS]):
        # xlrd 会整表载入内存，先按尺寸预判，超限直接跳过。
        if (
            worksheet.nrows > TMP_ATTACHMENT_TABLE_MAX_ROWS
            or worksheet.nrows * worksheet.ncols > TMP_ATTACHMENT_TABLE_MAX_CELLS
        ):
            continue
        try:
            rows = _bounded_table_rows(
                tuple(_clean_cell_value(worksheet.cell_value(row_number, col)) for col in range(worksheet.ncols))
                for row_number in range(worksheet.nrows)
            )
            sheets.append(
                _build_sheet_index(
                    name=worksheet.name,
                    sheet_index=sheet_index,
                    rows=rows,
                )
            )
        except ValueError:
            continue
    if not sheets:
        raise ValueError("表格中没有有效数据行")
    return {
        "file_name": file_name,
        "file_kind": "spreadsheet",
        "sheets": sheets,
        "sheets_truncated": workbook.nsheets > TMP_ATTACHMENT_TABLE_MAX_SHEETS,
    }


def _extract_csv_index(file_name: str, file_content: bytes) -> dict:
    try:
        text = file_content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = file_content.decode("gbk")

    rows = _bounded_table_rows(
        tuple(_clean_cell_value(cell) for cell in row) for row in csv.reader(io.StringIO(text))
    )
    sheet = _build_sheet_index(
        name=file_name,
        sheet_index=0,
        rows=rows,
        header_policy="first",
        allow_empty=True,
    )
    return {
        "file_name": file_name,
        "file_kind": "spreadsheet",
        "sheets": [sheet],
        "sheets_truncated": False,
    }


def _extract_table_index(file_name: str, file_content: bytes) -> dict:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".xlsx":
        return _extract_xlsx_index(file_name, file_content)
    if suffix == ".xls":
        return _extract_xls_index(file_name, file_content)
    if suffix == ".csv":
        return _extract_csv_index(file_name, file_content)
    raise ValueError(f"不支持的表格类型: {suffix}")


async def _prepare_tmp_parse_artifacts(
    *,
    file_name: str,
    file_content: bytes,
    bucket_name: str,
    object_name: str,
    tmp_file_id: str,
    uid: str,
) -> dict:
    """解析临时附件；统一生成 Markdown，表格额外生成结构化索引。"""
    suffix = Path(file_name).suffix.lower()
    markdown: str
    if suffix in TMP_ATTACHMENT_TABLE_EXTENSIONS:
        table_index = await asyncio.wait_for(
            asyncio.to_thread(_extract_table_index, file_name, file_content),
            timeout=TMP_ATTACHMENT_TABLE_SECONDS,
        )
        markdown = _build_table_markdown(table_index)
    elif suffix in TMP_ATTACHMENT_TEXT_EXTENSIONS:
        try:
            markdown = file_content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("文本文件不是有效的 UTF-8 内容") from exc
    else:
        markdown = await parse_document(_minio_source(bucket_name, object_name))

    markdown, truncated = _truncate_markdown(markdown)
    parsed_object_name = _make_tmp_parsed_object(str(uid), tmp_file_id, file_name)
    minio_client = get_minio_client()
    parsed_upload = await minio_client.aupload_file(
        bucket_name=bucket_name,
        object_name=parsed_object_name,
        data=markdown.encode("utf-8"),
        content_type="text/markdown; charset=utf-8",
    )
    result = {
        "parsed_object_name": parsed_upload.object_name,
        "parsed_minio_url": parsed_upload.url,
        "parse_method": "weknora",
        "truncated": truncated,
    }

    if suffix in TMP_ATTACHMENT_TABLE_EXTENSIONS:
        table_index_object_name = _make_tmp_table_index_object(str(uid), tmp_file_id, file_name)
        table_index_bytes = json.dumps(table_index, ensure_ascii=False).encode("utf-8")
        table_index_upload = await minio_client.aupload_file(
            bucket_name=bucket_name,
            object_name=table_index_object_name,
            data=table_index_bytes,
            content_type="application/json; charset=utf-8",
        )
        result["table_index_object_name"] = table_index_upload.object_name
        result["table_index_minio_url"] = table_index_upload.url

    return result


def _build_state_uploads(attachments: list[dict]) -> list[dict]:
    uploads: list[dict] = []
    for attachment in attachments:
        # 状态注入只信任请求快照；快照建立后不再读取可能被删除流程影响的附件记录。
        snapshot = attachment.get("request_snapshot")
        record = dict(snapshot) if isinstance(snapshot, dict) else dict(attachment)
        path = record.get("path")
        if not isinstance(path, str) or not path.strip():
            continue

        uploads.append(
            {
                "file_id": record.get("file_id"),
                "file_name": record.get("file_name"),
                "file_type": record.get("file_type"),
                "file_size": record.get("file_size", 0),
                "status": record.get("status", "uploaded"),
                "uploaded_at": record.get("uploaded_at"),
                "thread_id": record.get("thread_id"),
                "path": path,
                "original_path": record.get("original_path"),
                "table_index_path": record.get("table_index_path"),
                "artifact_url": record.get("artifact_url"),
                "request_id": record.get("request_id"),
                "parse_status": record.get("parse_status"),
                "parse_error": record.get("parse_error"),
                "truncated": bool(record.get("truncated")),
            }
        )
    return uploads


def serialize_attachment(record: dict) -> dict:
    snapshot = record.get("request_snapshot")
    source = dict(snapshot) if isinstance(snapshot, dict) else dict(record)
    path = source.get("path")
    return {
        "file_id": source.get("file_id"),
        "file_name": source.get("file_name"),
        "file_type": source.get("file_type"),
        "file_size": source.get("file_size", 0),
        "status": source.get("status", "uploaded"),
        "parse_status": source.get("parse_status"),
        "uploaded_at": source.get("uploaded_at"),
        "thread_id": source.get("thread_id"),
        "path": path,
        "artifact_url": source.get("artifact_url"),
        "original_path": source.get("original_path"),
        "original_artifact_url": source.get("original_artifact_url"),
        "source_bucket_name": source.get("source_bucket_name"),
        "source_object_name": source.get("source_object_name"),
        "source_minio_url": source.get("source_minio_url"),
        "minio_url": source.get("minio_url"),
        "request_id": source.get("request_id"),
        "parse_error": source.get("parse_error"),
        "truncated": bool(source.get("truncated")),
        "table_index_path": source.get("table_index_path"),
        "table_index_artifact_url": source.get("table_index_artifact_url"),
    }


def _materialize_tmp_attachment_files(
    *,
    thread_id: str,
    uid: str,
    file_id: str,
    file_name: str,
    file_content: bytes,
    parse_supported: bool,
    source_bucket_name: str,
    source_object_name: str,
    source_minio_url: str | None,
    parsed_markdown: str | None = None,
    truncated: bool = False,
    table_index: dict | None = None,
    parse_error: str | None = None,
) -> dict:
    """将 tmp 附件复制到线程目录，不主动删除 tmp 对象。"""
    ensure_thread_dirs(thread_id, uid)

    storage_name = f"{file_id}_{file_name}"
    upload_virtual_path = _make_upload_virtual_path(storage_name)
    uploads_dir = sandbox_uploads_dir(thread_id)
    upload_actual_path = uploads_dir / Path(upload_virtual_path).name
    upload_actual_path.write_bytes(file_content)

    record = {
        "thread_id": thread_id,
        "status": "uploaded",
        "parse_status": "unsupported" if not parse_supported else ("parse_failed" if parse_error else "parsed"),
        "parse_error": parse_error,
        "path": upload_virtual_path,
        "artifact_url": _artifact_url(thread_id, upload_virtual_path),
        "storage_path": str(upload_actual_path),
        "original_path": upload_virtual_path,
        "original_artifact_url": _artifact_url(thread_id, upload_virtual_path),
        "original_storage_path": str(upload_actual_path),
        "source_bucket_name": source_bucket_name,
        "source_object_name": source_object_name,
        "source_minio_url": source_minio_url,
        "minio_url": None,
    }

    if parsed_markdown is None:
        return record

    markdown_virtual_path, markdown_host_path = _build_attachment_storage_path(
        uid=uid,
        thread_id=thread_id,
        file_name=storage_name,
    )
    markdown_host_path.write_text(parsed_markdown, encoding="utf-8")
    record.update(
        {
            "status": "parsed",
            "parse_status": "parsed",
            "path": markdown_virtual_path,
            "artifact_url": _artifact_url(thread_id, markdown_virtual_path),
            "storage_path": str(markdown_host_path),
            "file_path": markdown_virtual_path,
            "markdown": parsed_markdown,
            "truncated": truncated,
            "markdown_storage_path": str(markdown_host_path),
        }
    )

    if table_index is not None:
        index_name = f"{Path(storage_name).stem}.index.json"
        index_virtual_path = f"{VIRTUAL_PATH_UPLOADS}/attachments/{index_name}"
        index_host_path = markdown_host_path.parent / index_name
        index_host_path.write_text(json.dumps(table_index, ensure_ascii=False), encoding="utf-8")
        record.update(
            {
                "table_index_path": index_virtual_path,
                "table_index_artifact_url": _artifact_url(thread_id, index_virtual_path),
                "table_index_storage_path": str(index_host_path),
            }
        )
    return record


async def create_thread_view(
    *,
    agent_slug: str,
    title: str | None,
    metadata: dict | None,
    db: AsyncSession,
    current_uid: str,
) -> dict:
    user_result = await db.execute(select(User).where(User.uid == str(current_uid)))
    current_user = user_result.scalar_one_or_none()
    if not current_user:
        raise HTTPException(status_code=404, detail="用户不存在")

    agent_repo = AgentRepository(db)
    agent_item = await agent_repo.get_visible_by_slug(slug=agent_slug, user=current_user)
    if not agent_item:
        raise HTTPException(status_code=404, detail="智能体不存在")

    thread_id = str(uuid.uuid4())
    conv_repo = ConversationRepository(db)
    thread_metadata = dict(metadata or {})
    thread_metadata["backend_id"] = agent_item.backend_id
    conversation = await conv_repo.create_conversation(
        uid=str(current_uid),
        agent_id=agent_item.slug,
        title=title or "新的对话",
        thread_id=thread_id,
        metadata=thread_metadata,
    )

    return {
        "id": conversation.thread_id,
        "uid": conversation.uid,
        "agent_id": conversation.agent_id,
        "title": conversation.title,
        "created_at": conversation.created_at.isoformat(),
        "updated_at": conversation.updated_at.isoformat(),
        "metadata": conversation.extra_metadata or {},
    }


async def list_threads_view(
    *,
    agent_slug: str | None,
    db: AsyncSession,
    current_uid: str,
    limit: int | None = None,
    offset: int = 0,
) -> list[dict]:
    conv_repo = ConversationRepository(db)
    conversations = await conv_repo.list_conversations(
        uid=str(current_uid),
        agent_id=agent_slug,
        status="active",
        limit=limit,
        offset=offset,
        exclude_sources=INVOCATION_CONVERSATION_SOURCES,
    )

    return [
        {
            "id": conv.thread_id,
            "uid": conv.uid,
            "agent_id": conv.agent_id,
            "title": conv.title,
            "is_pinned": bool(conv.is_pinned),
            "created_at": conv.created_at.isoformat(),
            "updated_at": conv.updated_at.isoformat(),
            "metadata": conv.extra_metadata or {},
        }
        for conv in conversations
    ]


async def delete_thread_view(
    *,
    thread_id: str,
    db: AsyncSession,
    current_uid: str,
) -> dict:
    conv_repo = ConversationRepository(db)
    await require_user_conversation(conv_repo, thread_id, str(current_uid))
    deleted = await conv_repo.delete_conversation(thread_id, soft_delete=True)
    if not deleted:
        raise HTTPException(status_code=404, detail="对话线程不存在")
    return {"message": "删除成功"}


async def update_thread_view(
    *,
    thread_id: str,
    title: str | None = None,
    is_pinned: bool | None = None,
    tool_approval_mode: str | None = None,
    db: AsyncSession,
    current_uid: str,
) -> dict:
    conv_repo = ConversationRepository(db)
    await require_user_conversation(conv_repo, thread_id, str(current_uid))
    metadata = {"tool_approval_mode": tool_approval_mode} if tool_approval_mode is not None else None
    updated_conv = await conv_repo.update_conversation(
        thread_id,
        title=title,
        is_pinned=is_pinned,
        metadata=metadata,
    )
    if not updated_conv:
        raise HTTPException(status_code=500, detail="更新失败")
    return {
        "id": updated_conv.thread_id,
        "uid": updated_conv.uid,
        "agent_id": updated_conv.agent_id,
        "title": updated_conv.title,
        "is_pinned": bool(updated_conv.is_pinned),
        "created_at": updated_conv.created_at.isoformat(),
        "updated_at": updated_conv.updated_at.isoformat(),
        "metadata": updated_conv.extra_metadata or {},
    }


async def upload_tmp_attachment_view(*, file: UploadFile, current_uid: str) -> dict:
    """上传附件到用户隔离的 MinIO tmp 路径。"""
    if not file.filename:
        raise HTTPException(status_code=400, detail="无法识别的文件名")

    file_name = _safe_file_name(file.filename)
    max_size_bytes = _get_attachment_max_size_bytes()
    try:
        file_content = await read_upload_with_limit(
            file,
            max_size_bytes=max_size_bytes,
            too_large_message=f"附件过大，当前仅支持 {max_size_bytes // (1024 * 1024)} MB 以内的文件",
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    file_size = len(file_content)
    tmp_file_id, object_name = _make_tmp_attachment_object(str(current_uid), file_name)
    minio_client = get_minio_client()
    bucket_name = _get_tmp_attachment_bucket()
    try:
        upload_result = await minio_client.aupload_file(
            bucket_name=bucket_name,
            object_name=object_name,
            data=file_content,
            content_type=file.content_type,
        )
    except StorageError as exc:
        raise HTTPException(status_code=500, detail=f"临时附件上传失败: {exc}") from exc

    suffix = Path(file_name).suffix.lower()
    parse_supported = suffix in TMP_ATTACHMENT_PARSE_EXTENSIONS
    parse_methods = list(TMP_ATTACHMENT_PARSE_METHODS) if parse_supported else []
    parse_status = "parsed" if parse_supported else "unsupported"
    parse_error = None
    parsed_payload: dict | None = None
    if parse_supported:
        try:
            parsed_payload = await _prepare_tmp_parse_artifacts(
                file_name=file_name,
                file_content=file_content,
                bucket_name=bucket_name,
                object_name=object_name,
                tmp_file_id=tmp_file_id,
                uid=str(current_uid),
            )
        except Exception as exc:  # noqa: BLE001
            # 上传和解析解耦：解析失败时保留原始附件，用户仍可选择重试或继续使用原文。
            parse_status = "parse_failed"
            parse_error = str(exc)
            logger.warning(f"Tmp attachment parse failed for {file_name}: {exc}")

    return {
        "tmp_file_id": tmp_file_id,
        "file_name": file_name,
        "file_type": file.content_type,
        "file_size": file_size,
        "bucket_name": upload_result.bucket_name,
        "object_name": upload_result.object_name,
        "minio_url": upload_result.url,
        "uploaded_at": utc_isoformat(),
        "parse_supported": bool(parse_methods),
        "parse_methods": parse_methods,
        "parse_status": parse_status,
        "parse_error": parse_error,
        "parsed_object_name": (parsed_payload or {}).get("parsed_object_name"),
        "parsed_minio_url": (parsed_payload or {}).get("parsed_minio_url"),
        "truncated": bool((parsed_payload or {}).get("truncated")),
        "table_index_object_name": (parsed_payload or {}).get("table_index_object_name"),
        "table_index_minio_url": (parsed_payload or {}).get("table_index_minio_url"),
    }


async def parse_tmp_attachment_view(
    *,
    object_name: str,
    file_name: str,
    parse_method: str | None,
    bucket_name: str | None,
    current_uid: str,
) -> dict:
    """解析用户 tmp 附件并把 markdown 写回 tmp。"""
    minio_client = get_minio_client()
    expected_bucket = _get_tmp_attachment_bucket()
    bucket_name = bucket_name or expected_bucket
    if not _is_allowed_tmp_attachment_bucket(bucket_name, expected_bucket):
        raise HTTPException(status_code=400, detail="无效的临时附件 bucket")

    tmp_file_id, safe_name = _require_tmp_object_section(object_name, str(current_uid), "original")
    method = _normalize_parse_method(safe_name, parse_method)

    try:
        parsed_result = await _prepare_tmp_parse_artifacts(
            file_name=safe_name,
            file_content=await minio_client.adownload_file(bucket_name, object_name),
            bucket_name=bucket_name,
            object_name=object_name,
            tmp_file_id=tmp_file_id,
            uid=str(current_uid),
        )
    except StorageError as exc:
        raise HTTPException(status_code=400, detail=f"读取临时附件失败: {exc}") from exc
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"Tmp attachment parse failed for {safe_name}: {exc}")
        raise HTTPException(status_code=400, detail=f"附件解析失败: {exc}") from exc

    return {
        "tmp_file_id": tmp_file_id,
        "file_name": safe_name,
        "bucket_name": bucket_name,
        "object_name": object_name,
        "parsed_object_name": parsed_result["parsed_object_name"],
        "parsed_minio_url": parsed_result["parsed_minio_url"],
        "parse_method": method,
        "status": "parsed",
        "truncated": parsed_result["truncated"],
        "table_index_object_name": parsed_result.get("table_index_object_name"),
        "table_index_minio_url": parsed_result.get("table_index_minio_url"),
    }


async def confirm_tmp_thread_attachments_view(
    *,
    thread_id: str,
    attachments: list[dict],
    db: AsyncSession,
    current_uid: str,
) -> dict:
    """将选中的 tmp 附件正式关联到对话线程。"""
    if not attachments:
        raise HTTPException(status_code=400, detail="请选择要添加的附件")

    conv_repo = ConversationRepository(db)
    conversation = await require_user_conversation(conv_repo, thread_id, str(current_uid))
    minio_client = get_minio_client()
    expected_bucket = _get_tmp_attachment_bucket()
    prepared_items: list[dict] = []

    for item in attachments:
        object_name = str(item.get("object_name") or "")
        bucket_name = str(item.get("bucket_name") or expected_bucket)
        if not _is_allowed_tmp_attachment_bucket(bucket_name, expected_bucket):
            raise HTTPException(status_code=400, detail="无效的临时附件 bucket")

        tmp_file_id, file_name = _require_tmp_object_section(object_name, str(current_uid), "original")
        try:
            file_content = await minio_client.adownload_file(bucket_name, object_name)
        except StorageError as exc:
            raise HTTPException(status_code=400, detail=f"读取临时附件失败: {exc}") from exc

        max_size_bytes = _get_attachment_max_size_bytes()
        if len(file_content) > max_size_bytes:
            max_size_mb = max_size_bytes // (1024 * 1024)
            raise HTTPException(status_code=400, detail=f"附件过大，当前仅支持 {max_size_mb} MB 以内的文件")

        parsed_markdown = None
        table_index = None
        parsed_object_name = str(item.get("parsed_object_name") or "")
        raw_parse_status = str(item.get("parse_status") or ("parsed" if parsed_object_name else "unsupported"))
        if raw_parse_status not in {"unsupported", "parse_failed", "parsed"}:
            raise HTTPException(status_code=400, detail="无效的附件解析状态")
        if raw_parse_status == "parsed" and not parsed_object_name:
            raise HTTPException(status_code=400, detail="解析状态为 parsed 时缺少解析附件")
        if raw_parse_status in {"unsupported", "parse_failed"} and parsed_object_name:
            raise HTTPException(status_code=400, detail="当前解析状态不应包含解析附件")
        if raw_parse_status == "parsed":
            _require_tmp_object_section(parsed_object_name, str(current_uid), "parsed", tmp_file_id)
            expected_parsed_object = _make_tmp_parsed_object(str(current_uid), tmp_file_id, file_name)
            if parsed_object_name != expected_parsed_object:
                raise HTTPException(status_code=400, detail="解析附件路径无效")
            try:
                parsed_bytes = await minio_client.adownload_file(bucket_name, parsed_object_name)
                parsed_markdown = parsed_bytes.decode("utf-8")
            except StorageError as exc:
                raise HTTPException(status_code=400, detail=f"读取解析附件失败: {exc}") from exc
            except UnicodeDecodeError as exc:
                raise HTTPException(status_code=400, detail="解析附件内容不是有效的 Markdown 文本") from exc
        parse_error = item.get("parse_error")
        if raw_parse_status == "parse_failed" and not isinstance(parse_error, str):
            parse_error = "附件解析失败"
        if raw_parse_status != "parse_failed":
            parse_error = None

        table_index_object_name = str(item.get("table_index_object_name") or "")
        if table_index_object_name:
            index_tmp_file_id, _index_section, index_file_name = _parse_user_tmp_object(
                table_index_object_name, str(current_uid)
            )
            expected_index_object = _make_tmp_table_index_object(str(current_uid), tmp_file_id, file_name)
            if (
                index_tmp_file_id != tmp_file_id
                or Path(index_file_name).suffix.lower() != ".json"
                or table_index_object_name != expected_index_object
            ):
                raise HTTPException(status_code=400, detail="表格索引路径无效")
            try:
                table_index_bytes = await minio_client.adownload_file(bucket_name, table_index_object_name)
                table_index = json.loads(table_index_bytes.decode("utf-8"))
            except StorageError as exc:
                raise HTTPException(status_code=400, detail=f"读取表格索引失败: {exc}") from exc
            except (UnicodeDecodeError, json.JSONDecodeError) as exc:
                raise HTTPException(status_code=400, detail="表格索引内容无效") from exc
        if raw_parse_status != "parsed":
            table_index = None

        prepared_items.append(
            {
                "file_name": file_name,
                "file_type": item.get("file_type"),
                "file_content": file_content,
                "parse_supported": raw_parse_status != "unsupported",
                "parse_status": raw_parse_status,
                "parse_error": parse_error,
                "parsed_markdown": parsed_markdown,
                "truncated": bool(item.get("truncated")) if raw_parse_status == "parsed" else False,
                "table_index": table_index,
                "source_bucket_name": bucket_name,
                "source_object_name": object_name,
                "source_minio_url": _minio_source(bucket_name, object_name),
            }
        )

    added_records: list[dict] = []
    for prepared in prepared_items:
        file_id = uuid.uuid4().hex
        materialized = _materialize_tmp_attachment_files(
            thread_id=thread_id,
            uid=str(conversation.uid),
            file_id=file_id,
            file_name=prepared["file_name"],
            file_content=prepared["file_content"],
            parse_supported=prepared["parse_supported"],
            source_bucket_name=prepared["source_bucket_name"],
            source_object_name=prepared["source_object_name"],
            source_minio_url=prepared["source_minio_url"],
            parsed_markdown=prepared["parsed_markdown"],
            truncated=prepared["truncated"],
            table_index=prepared["table_index"],
            parse_error=prepared["parse_error"],
        )
        attachment_record = {
            "thread_id": thread_id,
            "file_id": file_id,
            "file_name": prepared["file_name"],
            "file_type": prepared["file_type"],
            "file_size": len(prepared["file_content"]),
            "status": materialized["status"],
            "parse_status": materialized["parse_status"],
            "uploaded_at": utc_isoformat(),
            "path": materialized["path"],
            "artifact_url": materialized["artifact_url"],
            "storage_path": materialized["storage_path"],
            "original_path": materialized["original_path"],
            "original_artifact_url": materialized["original_artifact_url"],
            "original_storage_path": materialized["original_storage_path"],
            "source_bucket_name": materialized["source_bucket_name"],
            "source_object_name": materialized["source_object_name"],
            "source_minio_url": materialized["source_minio_url"],
            "minio_url": materialized["minio_url"],
        }
        for optional_key in (
            "file_path",
            "markdown",
            "truncated",
            "markdown_storage_path",
            "table_index_path",
            "table_index_artifact_url",
            "table_index_storage_path",
            "parse_status",
            "parse_error",
        ):
            if optional_key in materialized:
                attachment_record[optional_key] = materialized[optional_key]
        added_records.append(attachment_record)

    await conv_repo.add_attachments(conversation.id, added_records)
    # 已确认成功后清理临时对象；失败路径保留对象便于重试和排查。
    for prepared in prepared_items:
        try:
            tmp_file_id, _section, _file_name = _parse_user_tmp_object(
                prepared["source_object_name"], str(current_uid)
            )
            await minio_client.adelete_objects_by_prefix(
                prepared["source_bucket_name"],
                _tmp_attachment_prefix(str(current_uid), tmp_file_id),
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning(f"Failed to delete tmp attachment objects: {exc}")
    await invalidate_mention_cache(thread_id)

    return {"attachments": [serialize_attachment(item) for item in added_records]}


async def list_thread_attachments_view(
    *,
    thread_id: str,
    db: AsyncSession,
    current_uid: str,
) -> dict:
    conv_repo = ConversationRepository(db)
    conversation = await require_user_conversation(conv_repo, thread_id, str(current_uid))
    attachments = await conv_repo.get_attachments(conversation.id)
    return {
        "attachments": [serialize_attachment(item) for item in attachments],
        "limits": {
            "allowed_extensions": sorted(ATTACHMENT_ALLOWED_EXTENSIONS),
            "max_size_bytes": _get_attachment_max_size_bytes(),
        },
    }


async def delete_thread_attachment_view(
    *,
    thread_id: str,
    file_id: str,
    db: AsyncSession,
    current_uid: str,
) -> dict:
    conv_repo = ConversationRepository(db)
    conversation = await require_user_conversation(conv_repo, thread_id, str(current_uid))

    existing_attachments = await conv_repo.get_attachments(conversation.id)
    target_attachment = next((item for item in existing_attachments if item.get("file_id") == file_id), None)

    removed = await conv_repo.remove_attachment(conversation.id, file_id)
    if not removed:
        raise HTTPException(status_code=404, detail="附件不存在或已被删除")

    if target_attachment:
        delete_candidates = {
            str(value).strip()
            for value in (
                target_attachment.get("storage_path"),
                target_attachment.get("original_storage_path"),
                target_attachment.get("markdown_storage_path"),
                target_attachment.get("table_index_storage_path"),
            )
            if isinstance(value, str) and value.strip()
        }
        for candidate in delete_candidates:
            try:
                file_path = Path(candidate)
                if file_path.exists():
                    file_path.unlink()
            except Exception as exc:  # noqa: BLE001
                logger.warning(f"Failed to remove attachment file {candidate}: {exc}")

    source_bucket = target_attachment.get("source_bucket_name") if target_attachment else None
    source_object = target_attachment.get("source_object_name") if target_attachment else None
    if isinstance(source_bucket, str) and isinstance(source_object, str):
        try:
            tmp_file_id, _section, _file_name = _parse_user_tmp_object(source_object, str(current_uid))
            await get_minio_client().adelete_objects_by_prefix(
                source_bucket,
                _tmp_attachment_prefix(str(current_uid), tmp_file_id),
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning(f"Failed to delete tmp attachment objects: {exc}")

    await invalidate_mention_cache(thread_id)

    return {"message": "附件已删除"}


async def get_thread_history_view(
    *,
    thread_id: str,
    current_uid: str,
    db: AsyncSession,
) -> dict:
    """获取对话历史消息，包含用户反馈状态"""
    conv_repo = ConversationRepository(db)
    conversation = await conv_repo.get_conversation_by_thread_id(thread_id)
    if not conversation or conversation.uid != str(current_uid) or conversation.status == "deleted":
        raise HTTPException(status_code=404, detail="对话线程不存在")

    messages = await conv_repo.get_messages_by_thread_id(thread_id)
    messages = [
        message
        for message in messages
        if not (message.role == "user" and message.delivery_status in {"queued", "cancelled", "rejected"})
    ]

    run_ids_in_messages = {msg.run_id for msg in messages if msg.run_id}
    run_created_at: dict[str, Any] = {}
    if run_ids_in_messages:
        run_result = await db.execute(
            select(AgentRun.id, AgentRun.created_at)
            .where(AgentRun.id.in_(run_ids_in_messages))
            .order_by(AgentRun.created_at.asc(), AgentRun.id.asc())
        )
        run_created_at = {run_id: created_at for run_id, created_at in run_result.all()}
    messages.sort(
        key=lambda message: (
            run_created_at.get(message.run_id) or message.created_at,
            0 if message.role == "user" else 1,
            message.created_at,
            message.id,
        )
    )
    message_request_ids = set()
    for msg in messages:
        request_id = (msg.extra_metadata or {}).get("request_id")
        if msg.role == "user" and request_id:
            message_request_ids.add(str(request_id))
    attachments_by_request_id: dict[str, list[dict]] = {}
    if message_request_ids:
        for attachment in await conv_repo.get_attachments(conversation.id):
            request_id = attachment.get("request_id")
            if not request_id or str(request_id) not in message_request_ids:
                continue
            attachments_by_request_id.setdefault(str(request_id), []).append(serialize_attachment(attachment))

    history: list[dict] = []
    role_type_map = {"user": "human", "assistant": "ai", "tool": "tool", "system": "system"}

    for msg in messages:
        extra_metadata = dict(msg.extra_metadata or {})
        request_id = extra_metadata.get("request_id")
        if msg.role == "user" and request_id and not extra_metadata.get("attachments"):
            extra_metadata["attachments"] = attachments_by_request_id.get(str(request_id), [])

        msg_dict = {
            "id": msg.id,
            "type": role_type_map.get(msg.role, msg.role),
            "content": msg.content,
            "created_at": msg.created_at.isoformat() if msg.created_at else None,
            "run_id": msg.run_id,
            "request_id": msg.request_id,
            "delivery_status": msg.delivery_status,
            "error_type": extra_metadata.get("error_type"),
            "error_message": extra_metadata.get("error_message"),
            "extra_metadata": extra_metadata,
            "message_type": msg.message_type,
            "image_content": msg.image_content,
        }

        if msg.tool_calls:
            msg_dict["tool_calls"] = [
                {
                    "id": tc.langgraph_tool_call_id or str(tc.id),
                    "name": tc.tool_name,
                    "function": {"name": tc.tool_name},
                    "args": tc.tool_input or {},
                    "tool_call_result": {"content": (tc.tool_output or "")} if tc.status == "success" else None,
                    "status": tc.status,
                    "error_message": tc.error_message,
                }
                for tc in msg.tool_calls
            ]

        history.append(msg_dict)

    logger.info(f"Loaded {len(history)} messages for thread {thread_id}")
    return {"history": history}
