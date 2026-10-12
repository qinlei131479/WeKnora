#!/usr/bin/env python3
"""报价单工具：结构化取表（extract）/ 批量候选（candidates）/ 生成报价单（fill）/ 勾稽校验（verify）。

子命令：
  extract     读取标准或历史清单，输出表头、行号与有效行，供匹配时引用真实 Sheet + 行号。
  candidates  一次性为多条需求给出 Top-N 候选源行，替代逐条检索。
  fill        校验 rows.json 并生成报价单 Excel（默认子命令）。
  verify      纯 Python 勾稽生成的报价单：逐行 E+F+G=H、合计=明细和，不需要 LibreOffice。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

ALLOWED_RESULTS = {"已匹配", "未匹配"}
EFFORT_KEYS = ("product", "dev", "delivery")
EFFORT_LABELS = {"product": "产品", "dev": "研发", "delivery": "交付"}
EFFORT_LABEL_RE = re.compile(r"(产品|研发|交付)\s*=\s*(\d+(?:\.\d+)?)")
EFFORT_ITEM_RE = re.compile(
    r"「(?P<name>[^」]+)」\s*(?P<value>\d+(?:\.\d+)?)\s*（(?P<evidence>[^）]+)）"
)
LEVELS = ("L1", "L2", "L3", "L4", "T0", "T1", "T2", "T3")
UNMATCHED_PREFIX = "未从标准产品功能清单和历史项目库中找到"
UNMATCHED_SUFFIX = "匹配项"
ZERO_HINTS = ("无需", "不需要", "不涉及", "不投入", "已包含", "已具备", "复用", "规则0", "无历史参照")
HEADERS = [
    "一级功能",
    "二级功能",
    "三级功能",
    "核心功能描述",
    "产品时效",
    "研发时效",
    "交付时效",
    "时效总计",
    "工时说明",
    "功能匹配结果",
    "匹配模块/功能",
    "匹配依据说明",
    "匹配来源",
]
COL_WIDTHS = {
    "A": 14.37,
    "B": 16.13,
    "C": 14.63,
    "D": 48.0,
    "E": 14.37,
    "F": 12.0,
    "G": 16.37,
    "H": 12.0,
    "I": 48.0,
    "J": 15.25,
    "K": 15.75,
    "L": 33.38,
    "M": 26.12,
}
CJK_RE = re.compile(r"[\u4e00-\u9fff]+")
WORD_RE = re.compile(r"[A-Za-z0-9_]+")


# --------------------------------------------------------------------------- 通用工具


def _font(bold=False, size=11, name=None):
    kw = {"bold": bold, "size": size, "color": "FF0F1115"}
    if name:
        kw["name"] = name
    return Font(**kw)


def _data_font():
    return Font(size=11, color="FF0F1115")


def _thin_bottom():
    return Border(bottom=Side(style="thin"))


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_reason(text: str) -> str:
    """把半角括号、引号和逗号归一化为全角，模型写法差异不该导致出单失败。"""
    if not text:
        return text
    out = []
    quote_open = True
    for ch in text:
        if ch == "(":
            out.append("（")
        elif ch == ")":
            out.append("）")
        elif ch == ",":
            out.append("，")
        elif ch == '"' or ch == "'" or ch == "＂":
            out.append("「" if quote_open else "」")
            quote_open = not quote_open
        else:
            out.append(ch)
    return "".join(out)


def normalize_l3(text: str) -> str:
    """依赖标记统一为名称后紧跟 `*`，去掉多余空格、兼容全角＊。"""
    value = str(text or "").replace("＊", "*").strip()
    value = re.sub(r"\s+\*$", "*", value)
    value = re.sub(r"\*\s+", "*", value)
    return value


def _numeric_effort(row: dict, key: str) -> float:
    value = row.get(key)
    if isinstance(value, bool) or value is None:
        raise ValueError(key)
    if isinstance(value, str):
        text = value.strip()
        if text in ("", "--", "-", "—"):
            raise ValueError(key)
        value = text
    return float(value)


# --------------------------------------------------------------------------- 校验


def _validate_effort_reason(row: dict, i: int) -> list[str]:
    errors: list[str] = []
    reason = normalize_reason(str(row.get("effort_reason", "")).strip())
    matches = list(EFFORT_LABEL_RE.finditer(reason))
    if [m.group(1) for m in matches] != ["产品", "研发", "交付"]:
        return [
            f"第{i}行工时说明必须按 产品=数值：…；研发=数值：…；交付=数值：… 逐列写取数式；"
            f"三列都必须给出人天数字，不允许 --"
        ]

    for index, match in enumerate(matches):
        label = match.group(1)
        key = next(k for k, v in EFFORT_LABELS.items() if v == label)
        declared = float(match.group(2))
        segment = reason[match.start() : matches[index + 1].start() if index + 1 < len(matches) else len(reason)]

        try:
            row_value = _numeric_effort(row, key)
        except (TypeError, ValueError):
            errors.append(f"第{i}行{label}工时必须填人天数字（不允许 --），确实无需投入填 0 并写明原因")
            continue
        if row_value < 0:
            errors.append(f"第{i}行{label}工时不能小于 0")
        if abs(row_value - declared) > 0.051:
            errors.append(f"第{i}行{label}工时说明数值 {declared} 与 {key} 字段 {row_value} 不一致")

        items = list(EFFORT_ITEM_RE.finditer(segment))
        if not items:
            if declared == 0 and any(hint in segment for hint in ZERO_HINTS):
                continue
            errors.append(
                f"第{i}行{label}必须列出「功能/原子任务」数值（来源或AI评估定级原因）；"
                f"若确实无需投入则填 0 并在取数式里写明原因"
            )
            continue

        item_total = 0.0
        for item in items:
            name = item.group("name")
            value = float(item.group("value"))
            evidence = item.group("evidence")
            item_total += value
            if "AI原子任务" in name or "AI评估" in evidence:
                if "AI评估" not in evidence or not any(level in evidence for level in LEVELS):
                    errors.append(f"第{i}行{label}AI取数项 {name} 必须写 AI评估、定级和原因")
                else:
                    cause = evidence.replace("AI评估", "")
                    for level in LEVELS:
                        cause = cause.replace(level, "")
                    cause = re.sub(r"[，,、；;（)）(]", "", cause).strip()
                    if not cause:
                        errors.append(f"第{i}行{label}AI取数项 {name} 缺少定级原因")
            elif "Sheet：" not in evidence or "行：" not in evidence:
                errors.append(f"第{i}行{label}取数项 {name} 必须写 Sheet 和行号")
        if abs(item_total - declared) > 0.051:
            errors.append(f"第{i}行{label}取数项之和 {item_total:g} 必须等于 {label} 工时 {declared:g}")

    return errors


def _validate_merge_groups(rows: list[dict]) -> list[str]:
    errors: list[str] = []
    groups: dict[str, list[int]] = {}
    for i, row in enumerate(rows, 1):
        group = row.get("merge_group")
        if group in (None, ""):
            continue
        if not isinstance(group, str):
            errors.append(f"第{i}行 merge_group 必须是非空字符串")
            continue
        groups.setdefault(group, []).append(i)

    for group, row_numbers in groups.items():
        if len(row_numbers) != 2:
            errors.append(f"拆分组 {group} 必须恰好有 2 行")
            continue
        if row_numbers[1] - row_numbers[0] != 1:
            errors.append(f"拆分组 {group} 的 2 行必须相邻")
            continue
        first, second = rows[row_numbers[0] - 1], rows[row_numbers[1] - 1]
        if first.get("l1") != second.get("l1") or first.get("l2") != second.get("l2"):
            errors.append(f"拆分组 {group} 的一级、二级功能必须一致")
        results = sorted(str(row.get("match_result", "")) for row in (first, second))
        if results != ["已匹配", "未匹配"]:
            errors.append(f"拆分组 {group} 必须分别包含已匹配和未匹配")
    return errors


def validate(rows: list[dict]) -> list[str]:
    errors: list[str] = []
    if not rows:
        return ["rows 为空"]

    for i, row in enumerate(rows, 1):
        for key in (
            "l1",
            "l2",
            "l3",
            "desc",
            "effort_reason",
            "match_result",
            "match_module",
            "match_reason",
            "match_source",
        ):
            if key not in row or row[key] in (None, ""):
                errors.append(f"第{i}行缺少 {key}")

        row["l3"] = normalize_l3(row.get("l3"))
        row["effort_reason"] = normalize_reason(str(row.get("effort_reason", "")))

        for key in EFFORT_KEYS:
            try:
                value = _numeric_effort(row, key)
            except (TypeError, ValueError):
                errors.append(
                    f"第{i}行 {key} 必须填人天数字：未匹配行按 P3 原子任务拆解给值，"
                    f"确实无需投入填 0 并在取数式写明原因，不允许 --"
                )
            else:
                if value < 0:
                    errors.append(f"第{i}行 {key} 不能小于 0")

        result = row.get("match_result")
        if result not in ALLOWED_RESULTS:
            errors.append(f"第{i}行 match_result 非法: {result!r}")

        if result == "未匹配":
            if str(row.get("match_module", "")).strip() != "-":
                errors.append(f"第{i}行未匹配时 match_module 必须为 '-'")
            reason = str(row.get("match_reason", "")).strip()
            if not (reason.startswith(UNMATCHED_PREFIX) and reason.endswith(UNMATCHED_SUFFIX)):
                errors.append(
                    f"第{i}行未匹配依据必须以「{UNMATCHED_PREFIX}」开头、以「{UNMATCHED_SUFFIX}」结尾"
                )
            if "AI评估" not in str(row.get("match_source", "")):
                errors.append(f"第{i}行未匹配来源应含 AI评估")
            effort_reason = str(row.get("effort_reason", ""))
            if not ("AI原子任务" in effort_reason or "AI评估" in effort_reason):
                errors.append(
                    f"第{i}行未匹配必须完成 P3 原子任务拆解：工时说明里逐列写 AI原子任务「…」数值（AI评估，定级，原因）"
                )

        if result == "已匹配":
            match_reason = str(row.get("match_reason", ""))
            if "已支持" not in match_reason or "复用" not in match_reason:
                errors.append(f"第{i}行已匹配依据必须写已支持和复用")

        errors.extend(_validate_effort_reason(row, i))

        match_source = str(row.get("match_source", ""))
        effort_reason = str(row.get("effort_reason", ""))
        if result == "已匹配" and "标准" in match_source and "三项缺失，按规则0" in effort_reason:
            try:
                values = [_numeric_effort(row, key) for key in EFFORT_KEYS]
            except (TypeError, ValueError):
                errors.append(f"第{i}行标准三项缺失按规则0时，三列时效必须是数字")
            else:
                if any(value != 0 for value in values):
                    errors.append(f"第{i}行标准三项缺失按规则0时，三列时效必须都是 0")
        if result == "已匹配" and "历史" in match_source:
            if not any(marker in effort_reason for marker in ("直读", "累加", "求和", "+")):
                errors.append(f"第{i}行历史工时说明必须写直读、累加或求和计算式")

    errors.extend(_validate_merge_groups(rows))
    return errors


# --------------------------------------------------------------------------- fill


def fill(output: Path, rows: list[dict]) -> dict:
    errors = validate(rows)
    if errors:
        return {"ok": False, "errors": errors}

    wb = Workbook()
    ws = wb.active
    ws.title = "报价单"
    # 打开时强制重算：不依赖 LibreOffice 预写缓存值，Excel/WPS 打开即得正确结果。
    wb.calculation.fullCalcOnLoad = True

    header_font = _font(bold=True)
    data_font = _data_font()
    total_font = Font(name="等线", size=10)
    wrap_center = Alignment(vertical="center", wrap_text=True)
    center = Alignment(vertical="center")
    bottom = _thin_bottom()

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, header)
        cell.font = header_font
        cell.alignment = wrap_center if col in (4, 9, 10, 11, 12, 13) else center
        cell.border = bottom

    n = len(rows)
    merge_ranges: list[tuple[int, int]] = []
    merge_start: tuple[str, int] | None = None
    for i, row in enumerate(rows, 2):
        group = row.get("merge_group")
        if isinstance(group, str) and group:
            if merge_start and merge_start[0] == group:
                merge_ranges.append((merge_start[1], i))
                merge_start = None
            else:
                merge_start = (group, i)
        else:
            merge_start = None

        values = [
            row["l1"],
            row["l2"],
            row["l3"],
            row["desc"],
            _numeric_effort(row, "product"),
            _numeric_effort(row, "dev"),
            _numeric_effort(row, "delivery"),
            f"=E{i}+F{i}+G{i}",
            row["effort_reason"],
            row["match_result"],
            row["match_module"],
            row["match_reason"],
            row["match_source"],
        ]
        desc_len = max(len(str(row["desc"])), len(str(row["effort_reason"])))
        ws.row_dimensions[i].height = 44 if desc_len < 40 else 58 if desc_len < 80 else 72
        for col, val in enumerate(values, 1):
            cell = ws.cell(i, col, val)
            cell.font = data_font
            cell.border = bottom
            cell.alignment = wrap_center if col in (4, 9, 10, 11, 12, 13) else center

    for start, end in merge_ranges:
        ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
        ws.merge_cells(start_row=start, start_column=2, end_row=end, end_column=2)

    last = n + 1
    total_row = n + 2
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    ws.cell(total_row, 1, "合计").font = total_font
    ws.cell(total_row, 1).alignment = center
    for col, letter in enumerate(["E", "F", "G", "H"], 5):
        cell = ws.cell(total_row, col, f"=SUM({letter}2:{letter}{last})")
        cell.font = total_font
        cell.alignment = center
    ws.row_dimensions[total_row].height = 16.5

    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    totals = {key: sum(_numeric_effort(row, key) for row in rows) for key in EFFORT_KEYS}
    return {
        "ok": True,
        "rows": n,
        "product": round(totals["product"], 3),
        "dev": round(totals["dev"], 3),
        "delivery": round(totals["delivery"], 3),
        "total": round(sum(totals.values()), 3),
        "output": str(output),
        "merged_groups": len(merge_ranges),
        "recalc": "fullCalcOnLoad（打开即重算，无需 LibreOffice）",
        "next": f"fill_quote.py verify --file {output} --rows <rows.json>",
    }


# --------------------------------------------------------------------------- verify


def verify(path: Path, rows: list[dict]) -> dict:
    errors: list[str] = []
    if not path.exists():
        return {"ok": False, "errors": [f"文件不存在：{path}"]}

    wb = load_workbook(path)
    ws = wb["报价单"] if "报价单" in wb.sheetnames else wb.active

    n = len(rows)
    expected = {key: 0.0 for key in EFFORT_KEYS}
    for i, row in enumerate(rows, 2):
        try:
            values = [_numeric_effort(row, key) for key in EFFORT_KEYS]
        except (TypeError, ValueError):
            errors.append(f"第{i - 1}行 rows.json 缺少有效工时数字")
            continue
        for key, value in zip(EFFORT_KEYS, values):
            expected[key] += value
        for col, key in enumerate(EFFORT_KEYS, 5):
            cell_value = ws.cell(i, col).value
            if not isinstance(cell_value, (int, float)) or abs(float(cell_value) - values[col - 5]) > 0.001:
                errors.append(f"第{i}行第{col}列应为 {values[col - 5]:g}，实际 {cell_value!r}")
        formula = str(ws.cell(i, 8).value or "")
        if formula.replace(" ", "") != f"=E{i}+F{i}+G{i}":
            errors.append(f"第{i}行时效总计公式应为 =E{i}+F{i}+G{i}，实际 {formula!r}")
        row_total = sum(values)
        if abs(row_total - (values[0] + values[1] + values[2])) > 0.001:
            errors.append(f"第{i}行勾稽失败")

    if not wb.calculation.fullCalcOnLoad:
        errors.append("工作簿未设置 fullCalcOnLoad，打开时可能不重算")

    total_row = n + 2
    if _cell_text(ws.cell(total_row, 1).value) != "合计":
        errors.append(f"第{total_row}行应为合计行")
    for col, letter in enumerate(["E", "F", "G", "H"], 5):
        formula = str(ws.cell(total_row, col).value or "")
        if formula.replace(" ", "") != f"=SUM({letter}2:{letter}{n + 1})":
            errors.append(f"合计行第{col}列公式应为 =SUM({letter}2:{letter}{n + 1})，实际 {formula!r}")

    result: dict = {
        "ok": not errors,
        "file": str(path),
        "rows": n,
        "product": round(expected["product"], 3),
        "dev": round(expected["dev"], 3),
        "delivery": round(expected["delivery"], 3),
        "total": round(sum(expected.values()), 3),
        "checked": ["逐行 E/F/G 值", "逐行 H 公式", "合计行 SUM 公式", "fullCalcOnLoad"],
    }
    if errors:
        result["errors"] = errors
    return result


# --------------------------------------------------------------------------- extract / candidates


def _detect_header(grid: list[list]) -> int:
    for idx, raw in enumerate(grid[:20], start=1):
        if sum(1 for cell in raw if _cell_text(cell)) >= 3:
            return idx
    return 1


def _read_grid(path: Path, sheet: str | None) -> tuple[list[str], list[list], int, str]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets = list(wb.sheetnames)
        sheet_name = sheet or sheets[0]
        if sheet_name not in sheets:
            raise KeyError(f"Sheet 不存在：{sheet_name}，可用：{sheets}")
        ws = wb[sheet_name]
        grid = [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return sheets, grid, _detect_header(grid), sheet_name


def _sheets_of(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def extract(path: Path, sheet: str | None, max_rows: int, columns: list[str] | None) -> dict:
    if not path.exists():
        return {"ok": False, "errors": [f"文件不存在：{path}"]}
    grid = None
    try:
        _sheets, grid, header_row, sheet_name = _read_grid(path, sheet)
    except KeyError as exc:
        return {"ok": False, "errors": [str(exc)], "sheets": _sheets_of(path)}

    headers = [_cell_text(cell) for cell in grid[header_row - 1]] if grid else []
    picked: list[int] = []
    if columns:
        wanted = {c.strip() for c in columns if c.strip()}
        picked = [idx for idx, name in enumerate(headers) if name in wanted]
    if not picked:
        picked = list(range(len(headers)))

    rows: list[dict] = []
    for offset, raw in enumerate(grid[header_row:], start=header_row + 1):
        if len(rows) >= max_rows:
            break
        cells = {headers[idx]: _cell_text(raw[idx]) if idx < len(raw) else "" for idx in picked}
        if not any(cells.values()):
            continue
        rows.append({"row": offset, "cells": cells})

    return {
        "ok": True,
        "file": str(path),
        "sheet": sheet_name,
        "header_row": header_row,
        "headers": headers,
        "columns": [headers[idx] for idx in picked],
        "row_count": len(rows),
        "rows": rows,
    }


def _tokens(text: str) -> set[str]:
    text = str(text or "")
    tokens: set[str] = set()
    for run in CJK_RE.findall(text):
        if len(run) == 1:
            tokens.add(run)
        for idx in range(len(run) - 1):
            tokens.add(run[idx : idx + 2])
    tokens.update(word.lower() for word in WORD_RE.findall(text))
    return tokens


def _similarity(left: set[str], right: set[str]) -> float:
    if not left or not right:
        return 0.0
    inter = len(left & right)
    return inter / (len(left) * len(right)) ** 0.5


def _effort_columns(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, name in enumerate(headers):
        for key, label in EFFORT_LABELS.items():
            if key not in mapping and label in name:
                mapping[key] = idx
    return mapping


def candidates(path: Path, requirements: list[dict], sheet: str | None, top: int) -> dict:
    if not path.exists():
        return {"ok": False, "errors": [f"文件不存在：{path}"]}
    try:
        _sheets, grid, header_row, sheet_name = _read_grid(path, sheet)
    except KeyError as exc:
        return {"ok": False, "errors": [str(exc)], "sheets": _sheets_of(path)}

    headers = [_cell_text(cell) for cell in grid[header_row - 1]] if grid else []
    effort_cols = _effort_columns(headers)

    records: list[dict] = []
    for offset, raw in enumerate(grid[header_row:], start=header_row + 1):
        text = " ".join(_cell_text(cell) for cell in raw if cell is not None)
        if not text.strip():
            continue
        records.append({"row": offset, "raw": raw, "tokens": _tokens(text)})

    results = []
    for req in requirements:
        query = " ".join(
            str(req.get(field, "")) for field in ("l1", "l2", "l3", "desc", "requirement") if req.get(field)
        )
        query_tokens = _tokens(query)
        scored = sorted(
            ((_similarity(query_tokens, record["tokens"]), record) for record in records),
            key=lambda item: item[0],
            reverse=True,
        )
        picked = []
        for score, record in scored[:top]:
            picked.append(
                {
                    "row": record["row"],
                    "score": round(score, 4),
                    "cells": {headers[idx]: _cell_text(record["raw"][idx]) for idx in range(len(headers))},
                    "effort": {
                        key: _cell_text(record["raw"][idx]) if idx < len(record["raw"]) else ""
                        for key, idx in effort_cols.items()
                    },
                }
            )
        results.append(
            {
                "id": req.get("id"),
                "l3": req.get("l3"),
                "query": query,
                "candidates": picked,
            }
        )

    return {
        "ok": True,
        "file": str(path),
        "sheet": sheet_name,
        "header_row": header_row,
        "headers": headers,
        "effort_columns": {key: headers[idx] for key, idx in effort_cols.items()},
        "results": results,
    }


# --------------------------------------------------------------------------- CLI


def _load_rows(payload: dict | list) -> list[dict]:
    rows = payload["rows"] if isinstance(payload, dict) else payload
    if not isinstance(rows, list):
        raise ValueError("JSON 顶层必须是 rows 数组或数组本身")
    return rows


def _read_json(source: str) -> dict | list:
    raw = sys.stdin.read() if source == "-" else Path(source).read_text(encoding="utf-8")
    return json.loads(raw)


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv and not argv[0].startswith("-") and argv[0] not in ("extract", "candidates", "fill", "verify"):
        _die_unknown(argv[0])
        return 2
    if not argv or argv[0].startswith("-"):
        argv = ["fill", *argv]

    parser = argparse.ArgumentParser(description="报价单工具：extract / candidates / fill / verify")
    sub = parser.add_subparsers(dest="command")

    p_extract = sub.add_parser("extract", help="读取标准或历史清单，输出表头与行号")
    p_extract.add_argument("--file", required=True)
    p_extract.add_argument("--sheet", default=None)
    p_extract.add_argument("--max-rows", type=int, default=200)
    p_extract.add_argument("--columns", default=None, help="逗号分隔的列名，默认全列")

    p_cand = sub.add_parser("candidates", help="批量给出每条需求的 Top-N 候选源行")
    p_cand.add_argument("--file", required=True)
    p_cand.add_argument("--requirements", required=True, help="requirements JSON 路径，或 -")
    p_cand.add_argument("--sheet", default=None)
    p_cand.add_argument("--top", type=int, default=3)

    p_fill = sub.add_parser("fill", help="校验并生成报价单")
    p_fill.add_argument("--input", required=True, help="rows JSON 路径，或 -")
    p_fill.add_argument("--output", required=True)

    p_verify = sub.add_parser("verify", help="纯 Python 勾稽报价单，替代 recalc")
    p_verify.add_argument("--file", required=True)
    p_verify.add_argument("--rows", required=True, help="生成时使用的 rows JSON，或 -")

    args = parser.parse_args(argv)

    if args.command == "extract":
        columns = [c for c in args.columns.split(",")] if args.columns else None
        result = extract(Path(args.file), args.sheet, args.max_rows, columns)
    elif args.command == "candidates":
        payload = _read_json(args.requirements)
        reqs = payload["requirements"] if isinstance(payload, dict) else payload
        result = candidates(Path(args.file), reqs, args.sheet, args.top)
    elif args.command == "verify":
        rows = _load_rows(_read_json(args.rows))
        result = verify(Path(args.file), rows)
    else:
        rows = _load_rows(_read_json(args.input))
        result = fill(Path(args.output), rows)

    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result.get("ok") else 1


def _die_unknown(command: str) -> bool:
    print(
        json.dumps(
            {
                "ok": False,
                "errors": [f"未知子命令：{command}", "可用子命令：extract / candidates / fill / verify"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return True


if __name__ == "__main__":
    sys.exit(main())
